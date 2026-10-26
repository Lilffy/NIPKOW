# -*- coding: utf-8 -*-
"""
生成 A 的打印包（A4，DOCX）：
  1 封面/使用说明
  2 焊接顺序卡（7 步，打勾）
  3 上电流程卡
  4 测试点速查
  5 测试记录表
  6 采购清单（打勾）
  7 元件极性速查
  8-11 原理图分块图（4 张，各一页）
  12 一周学习打卡
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

WS = r"C:\旋转电视"
OUT = os.path.join(WS, "docs", "A_打印包.docx")
IMG = os.path.join(WS, "docs", "hardware", "sch_parts_print")

FONT_CN = "微软雅黑"
FONT_BODY = "宋体"


def set_cjk(run, name):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def body(doc, text, size=10.5, bold=False, space_after=4, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    set_cjk(r, FONT_BODY)
    if color:
        r.font.color.rgb = color
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(15)
    r.bold = True
    set_cjk(r, FONT_CN)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    # 下边框
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3A5F")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def add_runs(p, text, size, bold_all=False, font=FONT_BODY):
    """支持 **xxx** 语法 → 粗体片段"""
    import re as _re
    parts = _re.split(r"(\*\*.+?\*\*)", str(text))
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            r = p.add_run(part); r.bold = bold_all
        r.font.size = Pt(size)
        set_cjk(r, font)


def table(doc, headers, rows, widths=None, font_size=9.5, header_fill="DCE6F1"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # 固定布局，让列宽生效
    from docx.oxml import OxmlElement
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        set_cjk(r, FONT_CN)
        from docx.oxml import OxmlElement
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), header_fill)
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_runs(p, v, font_size)
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
            for cell in t.columns[i].cells:
                cell.width = Cm(w)
    return t


def page_break(doc):
    doc.add_page_break()


def main():
    doc = Document()

    # 页面：A4，页边距 1.6cm
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.orientation = WD_ORIENT.PORTRAIT
    for attr in ("top_margin", "bottom_margin"):
        setattr(s, attr, Cm(1.5))
    s.left_margin = Cm(1.6)
    s.right_margin = Cm(1.6)

    # 默认样式
    st = doc.styles["Normal"]
    st.font.name = FONT_BODY
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    # ============ P1 封面 ============
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run("NIPKOW 扩展板")
    r.font.size = Pt(30)
    r.bold = True
    set_cjk(r, FONT_CN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("A 的现场打印包")
    r.font.size = Pt(22)
    set_cjk(r, FONT_CN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("焊接 · 上电 · 测量 · 记录 · 采购")
    r.font.size = Pt(12)
    set_cjk(r, FONT_BODY)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    table(doc, ["本包包含", "打印建议"], [
        ["P2 焊接顺序卡（7 步打勾）", "A4 双面 / 单面都行，放工作台"],
        ["P3 上电流程卡", "单独一张，上电时拿在手上"],
        ["P4 测试点速查", "和 P3 一起"],
        ["P5 测试记录表", "多印 3–5 张（每次上电一张）"],
        ["P6 采购清单", "买料时带去，逐项划掉"],
        ["P7 元件极性速查", "焊接时对照"],
        ["P8–P11 原理图分块图", "建议彩印，A4"],
        ["P12 一周打卡 / 每日状态模板", "贴墙上"],
    ], widths=[9.0, 7.5], font_size=10)

    body(doc, "")
    body(doc, "版本：2026-10-26　对应 B 仓库 75533fc（C10 已改 10µF/25V，100µF 外接）", size=9,
         color=RGBColor(0x77, 0x77, 0x77))
    body(doc, "⚠️ 本包只服务于「焊接 + 上电 + 测量」，不做固件、不改原理图。", size=10, bold=True,
         color=RGBColor(0xB0, 0x30, 0x30))

    page_break(doc)

    # ============ P2 焊接顺序卡 ============
    h1(doc, "焊接顺序卡（每焊完一组就测一次）")
    body(doc, "原则：不要一次全焊完再上电。一组焊完 → 立刻测 → 通过 → 下一组。", size=10, bold=True)
    body(doc, "")

    rows = [
        ["1", "电源入口：J1、D1(SS34)、U2、U3、C1–C8、R10/R11/R16",
         "万用表通断档量 VIN–GND、5V–GND、3V3–GND", "不响（不短路）", "□"],
        ["2", "【不焊其他，先上电】9V / 限流 100–200mA",
         "量 TP1/TP2/TP3/TP4", "VIN≈9、5V≈5.0、3V3≈3.3，电流很小", "□"],
        ["3", "RGB 驱动：Q4/Q5/Q6(S8550) + R17–R22 + R3/R6/R9",
         "量三路集电极电压，或外接灯珠试亮", "三路一致，能点亮", "□"],
        ["4", "红外：R12、R13、R14、D2、J4、J5",
         "手挡红外对管，量 TP6 电平变化", "两态差 > 2V", "□"],
        ["5", "按键 KEY1/KEY2、JP1、J6、J7、R15",
         "按键按下量通断；J7 串口回环", "按下通、松开断", "□"],
        ["6", "两组 2×20 排母（H1/H2）",
         "目检 + 量相邻针脚", "不连锡、不歪", "□"],
        ["7", "测试点 TP1–TP9",
         "逐个别戳一下", "通到对应网络", "□"],
    ]
    table(doc, ["步", "焊什么", "焊完立刻做", "通过判据", "完成"],
          rows, widths=[1.0, 6.2, 5.6, 4.4, 1.3], font_size=9)

    body(doc, "")
    h1(doc, "不要焊（DNP）的器件")
    table(doc, ["位号", "器件", "为什么不贴"], [
        ["Q1、Q2、Q3", "AO3400A", "旧的低边方案，已改 PNP 高边"],
        ["R2、R5、R8", "10kΩ", "配 Q1–Q3 的栅极下拉"],
        ["R23", "0Ω", "J7-4 的 +3V3 预留，默认不贴"],
    ], widths=[4.0, 4.0, 9.0], font_size=9.5)
    body(doc, "其余按 BOM 全贴。", size=9.5)
    page_break(doc)

    # ============ P3 上电流程卡 ============
    h1(doc, "上电流程卡（最危险的一步）")
    body(doc, "记住一句话：限流电源 = 先把水龙头开到最小。", size=11, bold=True,
         color=RGBColor(0xB0, 0x30, 0x30))
    body(doc, "")

    table(doc, ["步", "做什么", "看到什么算正常", "不正常怎么办", "□"], [
        ["1", "不插核心板、不接电机、JP1 断开", "—", "—", "□"],
        ["2", "限流电源设 9V / 100mA，接 J1", "电流接近 0", "电流顶满 → 立刻断电，查短路", "□"],
        ["3", "量 TP1（VIN）", "9.0 ± 0.3 V", "不对 → 查 J1/D1 焊接", "□"],
        ["4", "量 TP2（VIN_PROTECTED）", "8.6–9.0 V（D1 压降约 0.4V）", "为 0 → D1 装反或虚焊", "□"],
        ["5", "量 TP3（+5V_REG）", "5.0 ± 0.2 V", "不对 → 查 U3/R10/R11", "□"],
        ["6", "量 TP4（+3V3_EXT）", "3.3 ± 0.1 V", "不对 → 查 U2", "□"],
        ["7", "摸 U3（TO-220）", "温温的；< 70°C", "烫手 → 断电，加散热片", "□"],
        ["8", "断电，插核心板，JP1 短接", "—", "—", "□"],
        ["9", "上电，量核心板 3.3V", "3.3V 稳定", "掉压 → 5V 带不动或短路", "□"],
        ["10", "接电机（不装鼓、不装桨）", "能转、能停", "不转 → 见固件/电平问题", "□"],
    ], widths=[1.0, 6.4, 5.2, 5.0, 1.2], font_size=9)

    body(doc, "")
    h1(doc, "三条保命规则")
    body(doc, "1. 电机第一次转：不装鼓、不装任何会飞的东西，手别伸进去。", size=10.5)
    body(doc, "2. 闻到糊味 / 听到“啪” / 板子烫手：先拔电，再看。", size=10.5)
    body(doc, "3. 改接线、动焊点之前：一定先断电。", size=10.5)
    body(doc, "")
    body(doc, "不要用 12V 适配器：U3 会多耗 0.7W，容易过热。用 9V。", size=10.5, bold=True,
         color=RGBColor(0xB0, 0x30, 0x30))
    page_break(doc)

    # ============ P4 测试点速查 ============
    h1(doc, "测试点速查表（TP1–TP9）")
    body(doc, "来源：B 的 A_HANDOVER_NOTES §5。以原理图为准，有出入在 issue 里指出。", size=9,
         color=RGBColor(0x66, 0x66, 0x66))
    body(doc, "")
    table(doc, ["测试点", "网络", "正常值", "怎么测"], [
        ["TP1", "VIN_MAIN", "9.0 V", "黑笔接 TP5(GND)，红笔戳 TP1"],
        ["TP2", "VIN_PROTECTED", "8.6–9.0 V", "同上（经过 D1 后）"],
        ["TP3", "+5V_REG", "5.0 ± 0.2 V", "同上"],
        ["TP4", "+3V3_EXT", "3.3 ± 0.1 V", "同上"],
        ["TP5", "GND", "0 V", "基准点，黑笔固定在这里"],
        ["TP6", "IR_INDEX (PA1)", "遮光 3.3V / 透光 ~0.3V", "手挡红外对管，看跳变"],
        ["TP7", "R_PWM (PA6)", "PWM 方波", "示波器；开漏+反相（低=亮）"],
        ["TP8", "G_PWM (PA7)", "PWM 方波", "同上"],
        ["TP9", "B_PWM (PB0) / MOTOR", "PWM 方波", "同上"],
    ], widths=[2.0, 5.2, 4.6, 6.0], font_size=9.5)
    body(doc, "")
    body(doc, "注：TP9 在部分版本里是电机 PWM，以丝印和原理图为准。", size=9, color=RGBColor(0x66, 0x66, 0x66))
    page_break(doc)

    # ============ P5 测试记录表 ============
    h1(doc, "测试记录表（每次上电填一张）")
    for _ in range(2):
        table(doc, ["项目", "记录", "项目", "记录"], [
            ["日期 / 板号", "", "限流设置", "____ V / ____ mA"],
            ["TP1 VIN", "____ V", "TP2 保护后", "____ V"],
            ["TP3 5V", "____ V", "TP4 3V3", "____ V"],
            ["U3 温度", "____ °C", "总电流", "____ mA"],
            ["RGB 红", "____ mA", "RGB 绿/蓝", "____ mA / ____ mA"],
            ["红外 遮光", "____ V", "红外 透光", "____ V"],
            ["三路 PWM", "有 / 无 波形", "频率/占空比", "____ / ____"],
            ["索引脉冲", "有 / 无，____ Hz", "电机 3 档", "____ / ____ / ____"],
            ["结论", "通过 / 不通过", "不通过原因", ""],
        ], widths=[3.6, 5.2, 3.6, 5.0], font_size=9.5)
        body(doc, "")
        body(doc, "签字：__________　　见证人：__________", size=9.5)
        body(doc, "")
    page_break(doc)

    # ============ P6 采购清单 ============
    h1(doc, "采购清单（按 B 的 hardware/BOM.csv）")
    body(doc, "阻容建议按 10 只起买（0603 太小，焊坏正常）。带 ★ 的是赛题/组委会提供的，不用买。", size=9.5)
    body(doc, "")
    table(doc, ["□", "位号", "器件", "规格", "立创编号", "数量"], [
        ["□", "D1", "肖特基", "SS34 / SMA", "C7420365", "1"],
        ["□", "D2", "LED", "GL0603UG01 / 0603", "C51933293", "1"],
        ["□", "U2", "LDO 3.3V", "LM1117-3.3 / SOT-223", "C55544533", "1"],
        ["★", "U3", "LDO 可调", "TO-220（组委会提供 C498319）", "—", "1"],
        ["★", "KEY1/2", "按键 5.8×5.8", "六脚插件（组委会提供）", "C5379898", "2"],
        ["□", "Q1–Q3", "AO3400A", "SOT-23（**DNP 不贴**）", "C700953", "0"],
        ["□", "Q4–Q6", "S8550", "SOT-23，丝印 2TY", "C20069126", "3"],
        ["□", "R3", "电阻 180Ω", "0603 1%", "通用", "1"],
        ["□", "R6/R9/R10", "电阻 120Ω", "0603 1%", "通用", "3"],
        ["□", "R11", "电阻 360Ω", "0603 1%", "通用", "1"],
        ["□", "R12", "电阻 200Ω", "0603 1%", "通用", "1"],
        ["□", "R13/R14/R17–R19", "电阻 1kΩ", "0603 1%", "通用", "5"],
        ["□", "R2/R5/R8/R15/R20–R22", "电阻 10kΩ", "0603 1%（R2/R5/R8 不贴）", "通用", "7"],
        ["□", "R16", "电阻 470Ω", "0603 1%", "通用", "1"],
        ["□", "R23", "0Ω", "0603（默认不贴）", "通用", "1"],
        ["□", "C1–C4", "10µF/16V", "0603 X5R/X7R（C1/C3 在 9V 上，**选 16V 以上**）", "通用", "4"],
        ["□", "C5–C9", "100nF/50V", "0603 X7R", "通用", "5"],
        ["□", "C10", "**10µF/25V**", "0603（已从 100µF 改）", "通用", "1"],
        ["□", "外接", "**100µF/25V 电解**", "直插，≥16V，接 J1 或 J6 电源端", "C72477", "1"],
        ["□", "J1", "电源端子", "XY301V-A-5.0-2P", "C9900016950", "1"],
        ["□", "J3", "XH2.54 4P", "A2547WV 立式", "C9900216472", "1"],
        ["□", "J4/J5", "排针 1×2", "2.54", "C9900015265", "2"],
        ["□", "J6", "排针 1×5", "MDF7-5P-2.54", "C19088617", "1"],
        ["□", "J7", "排针 1×4", "2.54", "C32713270", "1"],
        ["□", "JP1", "排针 1×2 + 短路帽", "2.54", "C9900015265", "1"],
        ["□", "TP1–TP9", "测试针", "RH-5000 Ø1.4", "C5277086", "9"],
        ["□", "H1/H2", "排母 2×20", "2.54 双排", "C53202170", "2"],
        ["□", "U1", "核心板", "天空星 STM32F407VET6（**组委会提供**）", "C22370035", "1"],
    ], widths=[0.9, 4.2, 3.6, 6.6, 2.4, 1.0], font_size=8.5)

    body(doc, "")
    body(doc, "工具：烙铁(可调温)、焊锡 0.8mm、助焊剂、吸锡带、尖头镊子、放大镜、万用表、可调限流电源、示波器",
         size=9.5)
    page_break(doc)

    # ============ P7 极性速查 ============
    h1(doc, "元件方向与极性速查")
    table(doc, ["器件", "怎么判断方向", "装反的后果"], [
        ["电解电容（外接 100µF）", "**长脚 = 正极**；外壳白边 = 负极", "鼓包、炸开"],
        ["SS34（D1）", "**白条端 = 负极（K）**，朝板内", "电源不通或短路"],
        ["LED（D2）", "短脚 = 负极；万用表二极管档正向有读数", "不亮"],
        ["S8550（Q4–Q6）", "丝印 `2TY`；按板上丝印方向（缺角/斜边对齐）", "灯不亮或常亮"],
        ["AO3400A（Q1–Q3）", "丝印 `A09T`；**不贴**", "—"],
        ["LM1117（U2）", "SOT-223：1=GND 2=OUT 3=IN，按丝印方向", "无输出/发热"],
        ["TO-220（U3）", "金属背板朝外（贴散热片），1=ADJ 2=OUT 3=IN", "无输出/发热"],
        ["排母 H1/H2", "1 脚方孔/有标记，别装反", "插核心板会烧"],
        ["按键 KEY1/2", "六脚，方向看引脚间距（按丝印外框）", "按了没反应"],
        ["排针 J4/J5/J7", "无极性", "—"],
    ], widths=[4.6, 9.0, 4.4], font_size=9.5)

    body(doc, "")
    body(doc, "焊接小技巧：先在一个焊盘点一点锡 → 镊子把元件推上去固定 → 再焊另一端。"
              "两个焊盘都先上锡，元件永远放不平。", size=9.5, bold=True)
    page_break(doc)

    # ============ P8–P11 分块图 ============
    figs = [
        ("01_核心板与排母.png", "① 核心板与排母（H1/H2）",
         "看：80 脚编号、40 脚 = +5V、38 脚 = +3V3、39/37 = GND。插核心板前必看。"),
        ("02_电源区.png", "② 电源区（J1 → D1 → U2/U3）",
         "看：D1 方向、U3 的 R10=120/R11=360 → 5V、C4/C8 在 5V 轨、JP1 把 5V 送进核心板。"),
        ("03_RGB驱动.png", "③ RGB 驱动（共阴，PNP 高边）",
         "看：Q4–Q6 发射极接 +5V、基极 1k + 10k 上拉、集电极经 R3/R6/R9 到 J3。Q1–Q3 不贴。"),
        ("04_红外与连接器.png", "④ 红外 / 按键 / 电机 / 串口",
         "看：R12=200Ω 发射、R13=1k 接收上拉、D2 指示、J6 五线、J7 串口。"),
    ]
    for i, (fn, title, note) in enumerate(figs):
        h1(doc, title)
        body(doc, note, size=10)
        path = os.path.join(IMG, fn)
        if os.path.exists(path):
            doc.add_picture(path, width=Cm(17.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            body(doc, f"[缺图：{fn}]", size=10, color=RGBColor(0xB0, 0x30, 0x30))
        if i < len(figs) - 1:
            page_break(doc)

    page_break(doc)

    # ============ P12 学习打卡 ============
    h1(doc, "一周学习打卡（等板期间）")
    table(doc, ["□", "天", "主题", "验收标准"], [
        ["□", "一", "认识元件", "抓 5 个元件说出方向/好坏"],
        ["□", "二", "万用表三档", "二极管 5 秒判断正负死活"],
        ["□", "三", "焊接练习", "废板焊 10 个 0603，8 个合格"],
        ["□", "四", "上电与限流", "能解释为什么先 100mA"],
        ["□", "五", "示波器", "测方波读出频率/幅度/占空比"],
        ["□", "六", "数据手册", "说清 9V→5V 为什么要散热"],
        ["□", "日", "记录与协作", "开一个带证据的 issue"],
    ], widths=[1.2, 1.4, 4.6, 10.8], font_size=10)

    body(doc, "")
    h1(doc, "每日状态模板（贴到 GitHub #18 或微信）")
    body(doc, "## YYYY-MM-DD  A", size=10.5, bold=True)
    body(doc, "- 做了：", size=10.5)
    body(doc, "- 下一步：", size=10.5)
    body(doc, "- 阻塞：", size=10.5)
    body(doc, "- 需要 B：", size=10.5)

    body(doc, "")
    body(doc, "原理图看着乱是正常的：B 用的是网络标签画法（同名网络自动相连）。"
              "认位置看 PCB 丝印，认值看 BOM，认方向看上面第 7 页。", size=9,
         color=RGBColor(0x66, 0x66, 0x66))

    doc.save(OUT)
    print("saved ->", OUT)


if __name__ == "__main__":
    main()
